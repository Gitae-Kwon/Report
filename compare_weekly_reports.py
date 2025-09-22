# compare_weekly_reports.py (요약)
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import difflib

def load_to_dataframe(file_or_path, sheet=None):
    name = getattr(file_or_path, "name", str(file_or_path)).lower()
    if name.endswith((".xlsx", ".xls")):
        return pd.read_excel(file_or_path, sheet_name=sheet)
    elif name.endswith(".pdf"):
        import pdfplumber
        frames = []
        with pdfplumber.open(file_or_path) as pdf:
            for page in pdf.pages:
                for tbl in (page.extract_tables() or []):
                    if tbl and len(tbl) >= 2:
                        header, rows = tbl[0], tbl[1:]
                        frames.append(pd.DataFrame(rows, columns=header))
        return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=["프로젝트명","런칭","금주 진행 업무"])
    else:
        raise ValueError("지원하지 않는 파일 형식입니다. (.xlsx, .xls, .pdf)")

def build_report(prev_df, curr_df, project_col, launch_col, work_col):
    merged = curr_df.merge(prev_df, on=project_col, how="outer", suffixes=("_curr","_prev"), indicator=True)
    def status_row(r):
        if r["_merge"]=="left_only": return "ADDED"
        if r["_merge"]=="right_only": return "REMOVED"
        if r[f"{launch_col}_curr"]!=r[f"{launch_col}_prev"] or r[f"{work_col}_curr"]!=r[f"{work_col}_prev"]:
            return "MODIFIED"
        return "UNCHANGED"
    merged["STATUS"]=merged.apply(status_row,axis=1)
    modified = merged[merged["STATUS"]=="MODIFIED"][[project_col,f"{launch_col}_prev",f"{launch_col}_curr",f"{work_col}_prev",f"{work_col}_curr"]]
    added    = merged[merged["STATUS"]=="ADDED"][[project_col,f"{launch_col}_curr",f"{work_col}_curr"]]
    removed  = merged[merged["STATUS"]=="REMOVED"][[project_col,f"{launch_col}_prev",f"{work_col}_prev"]]
    return merged, modified, added, removed

def write_excel(out_path, merged, modified, added, removed, project_col, launch_col, work_col):
    with pd.ExcelWriter(out_path, engine="openpyxl") as wr:
        merged.to_excel(wr, "Summary", index=False)
        modified.to_excel(wr, "Modified", index=False)
        added.to_excel(wr, "Added", index=False)
        removed.to_excel(wr, "Removed", index=False)
    wb = load_workbook(out_path)
    if "Modified" in wb.sheetnames:
        ws = wb["Modified"]
        yellow = PatternFill(start_color="FFF8B4", end_color="FFF8B4", fill_type="solid")
        for r in range(2, ws.max_row+1):
            ws.cell(row=r, column=3).fill = yellow
            ws.cell(row=r, column=5).fill = yellow
    wb.save(out_path)
